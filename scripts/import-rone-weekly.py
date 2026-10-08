#!/usr/bin/env python3
"""Download the official R-ONE weekly workbook and upload the latest 52 weeks."""

from __future__ import annotations

import json
import os
import shutil
import subprocess
import tempfile
import time
import urllib.parse
import urllib.request
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

RONE_BASE = "https://www.reb.or.kr/r-one/portal/bbs/rpt"
AUDIENCE = "yeongnam-property-atlas"
CHUNK_SIZE = 750

GROUPS = [
    ("인천광역시", "인천", "부산", {
        "제물포": "제물포구", "영종": "영종구", "미추홀": "미추홀구",
        "연수": "연수구", "남동": "남동구", "부평": "부평구",
        "계양": "계양구", "서해": "서해구", "검단": "검단구",
    }),
    ("부산광역시", "부산", "대구", {
        "중": "중구", "서": "서구", "동": "동구", "영도": "영도구",
        "부산진": "부산진구", "동래": "동래구", "남": "남구",
        "북": "북구", "해운대": "해운대구", "사하": "사하구",
        "금정": "금정구", "강서": "강서구", "연제": "연제구",
        "수영": "수영구", "사상": "사상구", "기장": "기장군",
    }),
    ("대구광역시", "대구", "대전", {
        "중": "중구", "동": "동구", "서": "서구", "남": "남구",
        "북": "북구", "수성": "수성구", "달서": "달서구", "달성": "달성군",
    }),
    ("울산광역시", "울산", "세종", {
        "중": "중구", "남": "남구", "동": "동구", "북": "북구", "울주": "울주군",
    }),
    ("경상북도", "경북", "경남", {
        "경주": "경주시", "구미": "구미시", "포항": "포항시",
        "김천": "김천시", "안동": "안동시", "영주": "영주시",
        "영천": "영천시", "상주": "상주시", "문경": "문경시",
        "경산": "경산시", "칠곡": "칠곡군",
    }),
    ("경상남도", "경남", "제주", {
        "창원": "창원시", "진주": "진주시", "통영": "통영시",
        "사천": "사천시", "김해": "김해시", "밀양": "밀양시",
        "거제": "거제시", "양산": "양산시",
    }),
]


def request_json(url: str, data: dict[str, str] | None = None) -> dict[str, Any]:
    encoded = urllib.parse.urlencode(data).encode() if data is not None else None
    request = urllib.request.Request(url, data=encoded, headers={"User-Agent": "property-atlas/1.0"})
    with urllib.request.urlopen(request, timeout=90) as response:
        return json.load(response)


def oidc_token() -> str:
    base = os.environ["ACTIONS_ID_TOKEN_REQUEST_URL"]
    separator = "&" if "?" in base else "?"
    request = urllib.request.Request(
        f"{base}{separator}audience={AUDIENCE}",
        headers={"Authorization": f"bearer {os.environ['ACTIONS_ID_TOKEN_REQUEST_TOKEN']}"},
    )
    with urllib.request.urlopen(request, timeout=30) as response:
        return str(json.load(response)["value"])


def post_rows(site_url: str, payload: dict[str, Any]) -> None:
    data = json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode()
    for attempt in range(3):
        try:
            completed = subprocess.run(
                [
                    "curl", "--fail-with-body", "--silent", "--show-error", "--retry", "2",
                    "-X", "POST", "-H", f"Authorization: Bearer {oidc_token()}",
                    "-H", "Content-Type: application/json", "--data-binary", "@-",
                    f"{site_url}/api/rone-weekly",
                ],
                input=data,
                capture_output=True,
                check=True,
                timeout=120,
            )
            result = json.loads(completed.stdout)
            if not result.get("ok"):
                raise RuntimeError(str(result))
            return
        except Exception as error:
            if attempt == 2:
                if isinstance(error, subprocess.CalledProcessError):
                    detail = (error.stderr or error.stdout or b"").decode(errors="replace")
                else:
                    detail = str(error)
                raise RuntimeError(f"사이트 주간자료 저장 실패: {detail[:500]}") from None
            time.sleep(2 ** attempt)


def find_workbook() -> tuple[Path, str]:
    listing = request_json(
        f"{RONE_BASE}/searchBulletin.do",
        {"page": "1", "rows": "20", "bbsCd": "RPT", "grpCd": "B1004", "listSubCd": "RPT03"},
    )
    bulletin = next(
        item for item in listing["data"]
        if item.get("bbsTit") == "주간아파트가격동향조사 시계열통계표"
    )
    seq = str(bulletin["seq"])
    detail = request_json(
        f"{RONE_BASE}/selectBulletin.do",
        {"bbsCd": "RPT", "listSubCd": "RPT03", "seq": seq, "noticeYn": "Y"},
    )["data"]
    attachment = next(item for item in detail["files"] if item.get("fileExt") == "xlsx")
    query = urllib.parse.urlencode({"bbsCd": "RPT", "seq": seq, "fileSeq": attachment["fileSeq"]})
    temporary = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    temporary.close()
    request = urllib.request.Request(
        f"{RONE_BASE}/downloadAttachFile.do?{query}",
        headers={"User-Agent": "property-atlas/1.0"},
    )
    with urllib.request.urlopen(request, timeout=180) as response, open(temporary.name, "wb") as output:
        shutil.copyfileobj(response, output)
    return Path(temporary.name), f"{attachment['viewFileNm']}.{attachment['fileExt']}"


def text(value: Any) -> str:
    return str(value or "").strip()


def target_columns(sheet: Any) -> list[tuple[int, str, str]]:
    row2 = [text(cell.value) for cell in sheet[2]]
    row3 = [text(cell.value) for cell in sheet[3]]
    targets: list[tuple[int, str, str]] = []
    for region, start_label, next_label, aliases in GROUPS:
        start = row2.index(start_label)
        end = row2.index(next_label, start + 1)
        targets.append((start, region, "전체"))
        parent = ""
        for index in range(start + 1, end):
            primary, secondary = row2[index], row3[index]
            if primary:
                parent = aliases.get(primary, "")
                district = aliases.get(primary)
            elif parent in {"포항시", "창원시"} and secondary:
                district = f"{parent} {secondary}구"
            else:
                district = None
            if district:
                targets.append((index, region, district))
    return targets


def workbook_rows(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sale_sheet = workbook["매매변동률"]
    dates = [
        value for (value,) in sale_sheet.iter_rows(min_col=1, max_col=1, values_only=True)
        if isinstance(value, datetime)
    ]
    latest = max(value.date() for value in dates)
    cutoff = sorted({value.date() for value in dates})[-52]
    rows: list[dict[str, Any]] = []
    for sheet_name, metric in (("매매변동률", "매매지수 주간"), ("전세변동률", "전세지수 주간")):
        sheet = workbook[sheet_name]
        targets = target_columns(sheet)
        maximum = max(index for index, _, _ in targets) + 1
        for values in sheet.iter_rows(min_row=6, max_col=maximum, values_only=True):
            raw_date = values[0]
            if not isinstance(raw_date, datetime) or raw_date.date() < cutoff:
                continue
            period = raw_date.date().isoformat()
            for index, region, district in targets:
                value = values[index]
                if isinstance(value, bool) or not isinstance(value, (int, float)):
                    continue
                rows.append({
                    "region": region,
                    "district": district,
                    "date": period,
                    "metric": metric,
                    "rate": round(float(value), 6),
                })
    workbook.close()
    return rows


def main() -> None:
    site_url = os.environ["SITE_URL"].rstrip("/")
    workbook_path, workbook_name = find_workbook()
    try:
        rows = workbook_rows(workbook_path)
    finally:
        workbook_path.unlink(missing_ok=True)
    if not rows:
        raise RuntimeError("R-ONE 주간 시계열 행을 찾지 못했습니다.")
    cutoff = min(str(row["date"]) for row in rows)
    verified_at = date.today().isoformat()
    for offset in range(0, len(rows), CHUNK_SIZE):
        chunk = rows[offset : offset + CHUNK_SIZE]
        final = offset + CHUNK_SIZE >= len(rows)
        post_rows(site_url, {
            "rows": chunk,
            "final": final,
            "total": len(rows),
            "workbook": workbook_name,
            "verifiedAt": verified_at,
            "cutoff": cutoff,
        })
        print(f"uploaded {min(offset + len(chunk), len(rows))}/{len(rows)}")


if __name__ == "__main__":
    main()
